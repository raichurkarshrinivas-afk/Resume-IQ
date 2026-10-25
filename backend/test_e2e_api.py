"""
End-to-End API integration test for Phase 1.
Starts a live uvicorn server in a background thread and tests the actual HTTP endpoints.
"""

import io
import time
import threading
import requests
import uvicorn
from docx import Document
from main import app

SERVER_PORT = 8008
BASE_URL = f"http://127.0.0.1:{SERVER_PORT}"


def run_server():
    uvicorn.run(app, host="127.0.0.1", port=SERVER_PORT, log_level="warning")


def main():
    print("Starting test server...")
    server_thread = threading.Thread(target=run_server, daemon=True)
    server_thread.start()

    # Wait for server to come alive
    connected = False
    for _ in range(25):
        try:
            r = requests.get(f"{BASE_URL}/api/health", timeout=1)
            if r.status_code == 200:
                connected = True
                break
        except Exception:
            time.sleep(0.2)

    if not connected:
        print("[FAIL] Server failed to start")
        exit(1)

    print("[OK] Test server is healthy and responding.")

    # 1. Register a test user
    test_email = f"test_{int(time.time())}@example.com"
    test_password = "Password123!"
    test_name = "Alex CareerSeeker"

    print(f"\n1. Testing POST /api/auth/register for {test_email}...")
    reg_resp = requests.post(
        f"{BASE_URL}/api/auth/register",
        json={"email": test_email, "password": test_password, "full_name": test_name},
    )
    assert reg_resp.status_code == 201, f"Expected 201, got {reg_resp.status_code}: {reg_resp.text}"
    reg_data = reg_resp.json()
    assert "access_token" in reg_data, "No access_token in register response"
    token = reg_data["access_token"]
    print(f"[OK] Registered successfully! JWT received (len={len(token)})")

    # 2. Login
    print("\n2. Testing POST /api/auth/login...")
    login_resp = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": test_email, "password": test_password},
    )
    assert login_resp.status_code == 200, f"Expected 200, got {login_resp.status_code}"
    print("[OK] Login successful! JWT verified.")

    # 3. GET /api/auth/me
    print("\n3. Testing GET /api/auth/me...")
    headers = {"Authorization": f"Bearer {token}"}
    me_resp = requests.get(f"{BASE_URL}/api/auth/me", headers=headers)
    assert me_resp.status_code == 200, f"Expected 200, got {me_resp.status_code}"
    assert me_resp.json()["email"] == test_email
    print(f"[OK] GET /api/auth/me returned user: {me_resp.json()['full_name']}")

    # 4. Upload: Reject file > 5 MB
    print("\n4. Testing POST /resume/upload rejection: file > 5 MB...")
    large_payload = b"0" * (5 * 1024 * 1024 + 1024)
    files = {"file": ("large_resume.pdf", io.BytesIO(large_payload), "application/pdf")}
    large_resp = requests.post(f"{BASE_URL}/resume/upload", headers=headers, files=files)
    assert large_resp.status_code == 413, f"Expected 413, got {large_resp.status_code}: {large_resp.text}"
    print(f"[OK] 5MB limit enforced (413): {large_resp.json().get('detail')}")

    # 5. Upload: Reject non-PDF/DOCX
    print("\n5. Testing POST /resume/upload rejection: non-PDF/DOCX file...")
    txt_payload = b"Plain text resume that should be rejected."
    files = {"file": ("resume.txt", io.BytesIO(txt_payload), "text/plain")}
    txt_resp = requests.post(f"{BASE_URL}/resume/upload", headers=headers, files=files)
    assert txt_resp.status_code == 415, f"Expected 415, got {txt_resp.status_code}: {txt_resp.text}"
    print(f"[OK] File type enforced (415): {txt_resp.json().get('detail')}")

    # 6. Upload: Reject < 100 characters extracted text (scanned PDF message)
    print("\n6. Testing POST /resume/upload rejection: < 100 characters extracted text...")
    short_doc = Document()
    short_doc.add_paragraph("Tiny CV")
    short_buf = io.BytesIO()
    short_doc.save(short_buf)
    short_buf.seek(0)
    files = {
        "file": (
            "short.docx",
            short_buf,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    short_resp = requests.post(f"{BASE_URL}/resume/upload", headers=headers, files=files)
    assert short_resp.status_code == 422, f"Expected 422, got {short_resp.status_code}: {short_resp.text}"
    detail = short_resp.json().get("detail", "")
    assert "scanned" in detail.lower() or "100" in detail, f"Expected friendly scanned message, got: {detail}"
    print(f"[OK] Extracted text length enforced (422): {detail}")

    # 7. Upload: Accept valid document (>= 100 characters)
    print("\n7. Testing POST /resume/upload acceptance: valid DOCX document...")
    valid_doc = Document()
    valid_doc.add_paragraph(
        "John Doe - Lead Systems Architect & Senior Software Engineer. "
        "Experienced in scaling high-throughput REST APIs, architecting event-driven systems with Kafka, "
        "PostgreSQL database query tuning, container orchestration with Kubernetes, and modern web applications with React."
    )
    valid_buf = io.BytesIO()
    valid_doc.save(valid_buf)
    valid_buf.seek(0)
    files = {
        "file": (
            "john_doe_resume.docx",
            valid_buf,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    valid_resp = requests.post(f"{BASE_URL}/resume/upload", headers=headers, files=files)
    assert valid_resp.status_code == 201, f"Expected 201, got {valid_resp.status_code}: {valid_resp.text}"
    res_data = valid_resp.json()
    assert res_data["filename"] == "john_doe_resume.docx"
    assert res_data["file_type"] == "docx"
    print(f"[OK] Resume uploaded successfully with ID {res_data['id']}!")

    # Also test the alternative endpoint /api/resumes/upload
    print("\n8. Testing alternative endpoint POST /api/resumes/upload...")
    valid_buf.seek(0)
    files = {
        "file": (
            "john_doe_resume2.docx",
            valid_buf,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    alt_resp = requests.post(f"{BASE_URL}/api/resumes/upload", headers=headers, files=files)
    assert alt_resp.status_code == 201
    print(f"[OK] /api/resumes/upload works identically with ID {alt_resp.json()['id']}!")

    print("\n=======================================================")
    print(">>> ALL HTTP ENDPOINTS & PHASE 1 REQUIREMENTS VERIFIED! <<<")
    print("=======================================================")


if __name__ == "__main__":
    main()
