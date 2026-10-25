"""
Content generation routes: cover letter, interview, roadmap, linkedin, bullets,
plus docx download, interview answer scoring, and roadmap progress tracking.
"""

import io
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from docx import Document
from core.dependencies import get_db, get_current_user
from models.user import User
from models.resume import Resume
from models.analysis import Analysis
from models.generated_content import GeneratedContent
from schemas.generated_content import (
    GenerateRequest,
    GeneratedContentOut,
    DownloadDocxRequest,
    InterviewScoreRequest,
    InterviewScoreResponse,
    RoadmapToggleRequest,
)
from services.generator import generate, score_candidate_answer

router = APIRouter(prefix="/api/generate", tags=["generate"])
alias_router = APIRouter(tags=["generate"])


def _get_user_analysis(analysis_id: int, db: Session, user: User) -> Analysis:
    """Helper: fetch analysis with ownership check, eagerly load resume."""
    analysis = (
        db.query(Analysis)
        .options(joinedload(Analysis.resume))
        .join(Resume)
        .filter(Analysis.id == analysis_id, Resume.user_id == user.id)
        .first()
    )
    if not analysis:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Analysis not found")
    return analysis


@router.post("/cover-letter", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
@alias_router.post("/generate/cover-letter", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
async def generate_cover_letter(
    body: GenerateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Generate a cover letter based on an analysis (250-300 words, factual only)."""
    analysis = _get_user_analysis(body.analysis_id, db, current_user)
    result = await generate(
        db=db,
        analysis=analysis,
        content_type="cover_letter",
        tone=body.tone or "Formal",
        company_name=body.company_name,
    )
    return GeneratedContentOut.model_validate(result)


@router.post("/interview-prep", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
@alias_router.post("/generate/interview-prep", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
async def generate_interview_prep(
    body: GenerateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Generate 10 structured interview questions (5 tech, 3 behavioral, 2 weak spots)."""
    analysis = _get_user_analysis(body.analysis_id, db, current_user)
    result = await generate(
        db=db,
        analysis=analysis,
        content_type="interview",
        num_questions=10,
    )
    return GeneratedContentOut.model_validate(result)


@router.post("/career-roadmap", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
@alias_router.post("/generate/career-roadmap", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
async def generate_career_roadmap(
    body: GenerateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Generate a 30/60/90-day learning roadmap addressing missing keywords."""
    analysis = _get_user_analysis(body.analysis_id, db, current_user)
    result = await generate(
        db=db,
        analysis=analysis,
        content_type="roadmap",
        target_role=body.target_role,
    )
    return GeneratedContentOut.model_validate(result)


@router.post("/linkedin-summary", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
@alias_router.post("/generate/linkedin-summary", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
async def generate_linkedin_summary(
    body: GenerateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Generate 3 headline options, About section under 2000 chars, and 15 skills."""
    analysis = _get_user_analysis(body.analysis_id, db, current_user)
    result = await generate(
        db=db,
        analysis=analysis,
        content_type="linkedin",
        tone=body.tone,
    )
    return GeneratedContentOut.model_validate(result)


@router.post("/bullet-points", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
@alias_router.post("/generate/bullet-points", response_model=GeneratedContentOut, status_code=status.HTTP_201_CREATED)
async def generate_bullet_points(
    body: GenerateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Generate improved resume bullet points using XYZ formula with [X%] placeholders."""
    analysis = _get_user_analysis(body.analysis_id, db, current_user)
    result = await generate(
        db=db,
        analysis=analysis,
        content_type="bullets",
        section=body.section,
    )
    return GeneratedContentOut.model_validate(result)


@router.post("/download-docx")
@alias_router.post("/download-docx")
def download_docx(body: DownloadDocxRequest):
    """Generate and stream a .docx file for cover letter download."""
    doc = Document()
    doc.add_heading(body.title, 0)
    for paragraph in body.text.split("\n\n"):
        trimmed = paragraph.strip()
        if trimmed:
            doc.add_paragraph(trimmed)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    filename = f"{body.title.replace(' ', '_')}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# Interview Answer Scoring: POST /interview/score & POST /api/interview/score
@alias_router.post("/interview/score", response_model=InterviewScoreResponse)
@alias_router.post("/api/interview/score", response_model=InterviewScoreResponse)
@router.post("/interview/score", response_model=InterviewScoreResponse)
def score_interview(body: InterviewScoreRequest):
    """Score candidate interview response out of 10 with short feedback badge."""
    return score_candidate_answer(
        question=body.question,
        answer=body.answer,
        job_title=body.job_title,
    )


# Roadmap Checkbox Progress: POST /api/analysis/{analysis_id}/roadmap-toggle
@alias_router.post("/api/analysis/{analysis_id}/roadmap-toggle")
@alias_router.post("/analysis/{analysis_id}/roadmap-toggle")
@router.post("/analysis/{analysis_id}/roadmap-toggle")
@alias_router.post("/roadmap/toggle")
def toggle_roadmap_item(
    analysis_id: int,
    body: RoadmapToggleRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Persist roadmap checkbox state across reloads."""
    analysis = (
        db.query(Analysis)
        .join(Resume)
        .filter(Analysis.id == analysis_id, Resume.user_id == current_user.id)
        .first()
    )
    if not analysis:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Analysis not found")

    res_json = dict(analysis.result_json or {})
    roadmap_progress = dict(res_json.get("roadmap_progress") or {})
    roadmap_progress[body.item_id] = body.completed
    res_json["roadmap_progress"] = roadmap_progress
    analysis.result_json = res_json

    db.commit()
    return {"status": "ok", "roadmap_progress": roadmap_progress}


@router.get("/history", response_model=list[GeneratedContentOut])
@alias_router.get("/generate/history", response_model=list[GeneratedContentOut])
def get_generation_history(
    analysis_id: int | None = None,
    content_type: str | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """List previously generated content, optionally filtered."""
    query = (
        db.query(GeneratedContent)
        .join(Analysis)
        .join(Resume)
        .filter(Resume.user_id == current_user.id)
    )
    if analysis_id is not None:
        query = query.filter(GeneratedContent.analysis_id == analysis_id)
    if content_type is not None:
        query = query.filter(GeneratedContent.content_type == content_type)

    results = query.order_by(GeneratedContent.created_at.desc()).limit(50).all()
    return [GeneratedContentOut.model_validate(r) for r in results]
