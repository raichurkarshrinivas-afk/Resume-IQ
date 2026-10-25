"""
Resume file parser — extracts raw text and document metadata from PDF and DOCX files.
"""

import io
import re
import pdfplumber
from docx import Document
from fastapi import HTTPException, UploadFile, status


ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


class ParseResult(tuple):
    """
    Backwards-compatible tuple subclass (raw_text, file_type) with metadata attribute.
    Supports unpacking: `raw_text, file_type = await parse_file(file)`
    as well as attribute access: `result.metadata`.
    """
    def __new__(cls, raw_text: str, file_type: str, metadata: dict | None = None):
        return super().__new__(cls, (raw_text, file_type))

    def __init__(self, raw_text: str, file_type: str, metadata: dict | None = None):
        self.raw_text = raw_text
        self.file_type = file_type
        self.metadata = metadata or {}


def _detect_sections(text_lower: str) -> dict[str, bool]:
    """Detect presence of standard ATS resume sections."""
    return {
        "experience": bool(re.search(r"\b(experience|employment|work history|career)\b", text_lower)),
        "skills": bool(re.search(r"\b(skills|technical skills|competencies|technologies|tools)\b", text_lower)),
        "education": bool(re.search(r"\b(education|degree|university|bachelor|master|b\.s|m\.s)\b", text_lower)),
        "summary": bool(re.search(r"\b(summary|objective|professional summary|profile|about me)\b", text_lower)),
        "projects": bool(re.search(r"\b(projects|key projects|portfolio|open source)\b", text_lower)),
    }


async def parse_file(file: UploadFile) -> ParseResult:
    """
    Extract text and layout metadata from an uploaded PDF or DOCX file.

    Returns:
        ParseResult tuple: (raw_text, file_type) with .metadata
    """
    # Validate content type
    content_type = file.content_type or ""
    file_type = ALLOWED_TYPES.get(content_type)

    if not file_type:
        # Fallback: check extension
        filename = file.filename or ""
        if filename.lower().endswith(".pdf"):
            file_type = "pdf"
        elif filename.lower().endswith(".docx"):
            file_type = "docx"
        else:
            raise HTTPException(
                status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
                detail="Only PDF and DOCX files are supported",
            )

    # Read file bytes
    content = await file.read()

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="File size exceeds the 5 MB limit",
        )

    # Extract text & metadata
    metadata = {
        "file_size_bytes": len(content),
        "file_type": file_type,
        "page_count": 1,
        "tables_detected": False,
        "has_standard_sections": {},
        "char_count": 0,
        "word_count": 0,
        "clean_encoding": True,
    }

    try:
        if file_type == "pdf":
            raw_text, pdf_meta = _extract_pdf(content)
            metadata.update(pdf_meta)
        else:
            raw_text, docx_meta = _extract_docx(content)
            metadata.update(docx_meta)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Unable to parse the {file_type.upper()} file. Please ensure the document is not password-protected or corrupted.",
        )

    cleaned_text = raw_text.strip()
    if len(cleaned_text) < 100:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "Could not extract sufficient text from your resume (less than 100 characters detected). "
                "If this is an image-based or scanned PDF, please provide a text-searchable PDF or DOCX file."
            ),
        )

    words = cleaned_text.split()
    section_map = _detect_sections(cleaned_text.lower())
    detected_list = [k for k, v in section_map.items() if v]

    metadata["char_count"] = len(cleaned_text)
    metadata["character_count"] = len(cleaned_text)
    metadata["word_count"] = len(words)
    metadata["has_standard_sections"] = any(section_map.values())
    metadata["section_map"] = section_map
    metadata["detected_sections"] = detected_list
    metadata["clean_encoding"] = "\ufffd" not in raw_text
    metadata["encoding_clean"] = "\ufffd" not in raw_text

    return ParseResult(raw_text, file_type, metadata)


def _extract_pdf(content: bytes) -> tuple[str, dict]:
    """Extract text and metadata from PDF bytes using pdfplumber."""
    text_parts = []
    has_tables = False
    page_count = 0

    with pdfplumber.open(io.BytesIO(content)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            tables = page.extract_tables()
            if tables:
                has_tables = True
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

    raw_text = "\n\n".join(text_parts)
    return raw_text, {
        "page_count": page_count,
        "tables_detected": has_tables,
    }


def _extract_docx(content: bytes) -> tuple[str, dict]:
    """Extract text and metadata from DOCX bytes using python-docx."""
    doc = Document(io.BytesIO(content))
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    raw_text = "\n".join(paragraphs)

    has_tables = len(doc.tables) > 0
    # Estimate pages (~400 words per page in typical resumes)
    word_count = len(raw_text.split())
    estimated_pages = max(1, round(word_count / 400))

    return raw_text, {
        "page_count": estimated_pages,
        "tables_detected": has_tables,
        "paragraph_count": len(paragraphs),
    }
