"""
Live HTTP Integration Test Suite for Phase 4 & Phase 5
Tests:
1. Cover letter generation & tone selection (Formal, Friendly, Confident strictly 250-300 words)
2. DOCX streaming endpoint (/api/generate/download-docx and /download-docx)
3. 10 Interview questions (5 tech, 3 behavioral, 2 weak-spot) with why_asked and STAR outlines
4. POST /interview/score and /api/interview/score
5. 30/60/90-day Roadmap & DB progress checkbox toggle
6. LinkedIn summary with 3 headlines, <2000 char About, 15 skills, checklist
7. Analysis history listing for comparison
"""

import io
import time
import threading
import requests
import uvicorn
from docx import Document
from main import app

SERVER_PORT = 8012
BASE_URL = f"http://127.0.0.1:{SERVER_PORT}"


def run_server():
    uvicorn.run(app, host="127.0.0.1", port=SERVER_PORT, log_level="warning")


def main():
    print(f"Starting test server on port {SERVER_PORT}...")
    server_thread = threading.Thread(target=run_server, daemon=True)
    server_thread.start()

    # Wait for server
    connected = False
    for _ in range(30):
        try:
            r = requests.get(f"{BASE_URL}/api/health", timeout=1)
            if r.status_code == 200:
                connected = True
                break
        except Exception:
            time.sleep(0.2)

    if not connected:
        print("[FAIL] Server failed to start on port", SERVER_PORT)
        exit(1)

    print("[OK] Test server is healthy and responding.")

    # 1. Register test user
    user_email = f"phase45_{int(time.time())}@example.com"
    user_pwd = "SecurePassword123!"

    reg_res = requests.post(f"{BASE_URL}/auth/register", json={
        "email": user_email,
        "password": user_pwd,
        "full_name": "Phase 45 Test Candidate"
    })
    assert reg_res.status_code in [200, 201], f"Register failed: {reg_res.text}"
    token = reg_res.json()["access_token"]
    auth_headers = {"Authorization": f"Bearer {token}"}

    # 2. Upload resume
    resume_text = (
        "Jane Doe - Staff Software Engineer\n"
        "Summary: Experienced engineer specializing in Python, FastAPI, Docker, and distributed systems architecture.\n"
        "Experience:\n"
        "Senior Backend Engineer at TechCorp (2020-2024)\n"
        "- Architected high-throughput microservices using Python, FastAPI, and PostgreSQL handling millions of events.\n"
        "- Automated CI/CD deployment pipelines using Docker and Kubernetes, reducing release times.\n"
        "- Refactored legacy monolithic services, improving query performance.\n"
        "Education:\n"
        "B.S. in Computer Science, State University\n"
        "Skills: Python, FastAPI, PostgreSQL, Docker, AWS, Distributed Systems, Git, CI/CD"
    )

    doc = Document()
    doc.add_heading("Jane Doe - Staff Software Engineer", 0)
    for line in resume_text.split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())
    doc_bytes = io.BytesIO()
    doc.save(doc_bytes)
    doc_bytes.seek(0)

    upload_res = requests.post(
        f"{BASE_URL}/resume/upload",
        headers=auth_headers,
        files={"file": ("resume.docx", doc_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"override_text": resume_text}
    )
    assert upload_res.status_code == 201, f"Upload failed: {upload_res.text}"
    resume_id = upload_res.json()["id"]

    # 3. Run analysis
    analysis_res = requests.post(
        f"{BASE_URL}/analyze",
        headers=auth_headers,
        json={
            "resume_id": resume_id,
            "job_title": "Lead Backend Engineer",
            "job_description": "Seeking Lead Backend Engineer skilled in Python, FastAPI, AWS, Redis, Kafka, and Kubernetes."
        }
    )
    assert analysis_res.status_code == 201, f"Analyze failed: {analysis_res.text}"
    analysis_id = analysis_res.json()["id"]

    # 4. Test Cover letter tones & 250-300 word counts
    print("[1/6] Testing Cover Letter tones and 250-300 word counts...")
    for tone in ["Formal", "Friendly", "Confident"]:
        res = requests.post(
            f"{BASE_URL}/api/generate/cover-letter",
            headers=auth_headers,
            json={
                "analysis_id": analysis_id,
                "tone": tone,
                "company_name": "Acme Innovations"
            }
        )
        assert res.status_code == 201, f"Cover letter generation failed: {res.text}"
        data = res.json()
        assert data["content_type"] == "cover_letter"
        words = data["content"].split()
        word_count = len(words)
        assert 250 <= word_count <= 300, f"Tone {tone} returned {word_count} words (expected 250-300)"
        assert "Acme Innovations" in data["content"]
        print(f"      - {tone}: {word_count} words [PASS]")
    print("      PASSED: Formal, Friendly, and Confident all strictly 250-300 words!")

    # 5. Test DOCX download streaming endpoint
    print("[2/6] Testing DOCX download streaming endpoint...")
    for endpoint in ["/api/generate/download-docx", "/download-docx"]:
        res = requests.post(
            f"{BASE_URL}{endpoint}",
            json={
                "title": "Cover_Letter_Test",
                "text": "Dear Hiring Manager,\n\nThis is paragraph one.\n\nThis is paragraph two."
            }
        )
        assert res.status_code == 200, f"DOCX download failed at {endpoint}: {res.status_code}"
        assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in res.headers["content-type"]
        bio = io.BytesIO(res.content)
        doc = Document(bio)
        assert len(doc.paragraphs) >= 2
    print("      PASSED: /api/generate/download-docx and /download-docx streamed valid .docx!")

    # 6. Test Interview prep 10 questions distribution
    print("[3/6] Testing Interview questions distribution (5 tech, 3 behav, 2 weak)...")
    res = requests.post(
        f"{BASE_URL}/api/generate/interview-prep",
        headers=auth_headers,
        json={"analysis_id": analysis_id}
    )
    assert res.status_code == 201
    import json
    questions = json.loads(res.json()["content"])
    assert len(questions) == 10

    tech = [q for q in questions if q["category"] == "Technical"]
    behav = [q for q in questions if q["category"] == "Behavioral"]
    weak = [q for q in questions if q["category"] == "Weak spot"]

    assert len(tech) == 5, f"Expected 5 tech, got {len(tech)}"
    assert len(behav) == 3, f"Expected 3 behavioral, got {len(behav)}"
    assert len(weak) == 2, f"Expected 2 weak spots, got {len(weak)}"

    for q in questions:
        assert "why_asked" in q
        assert "star_outline" in q
        star = q["star_outline"]
        assert "situation" in star and "task" in star and "action" in star and "result" in star
    print("      PASSED: Exactly 10 questions with categories, why_asked, and STAR outlines!")

    # 7. Test Interview answer scoring endpoint
    print("[4/6] Testing Interview answer scoring endpoint...")
    for endpoint in ["/interview/score", "/api/interview/score"]:
        res = requests.post(
            f"{BASE_URL}{endpoint}",
            json={
                "question": "How do you handle microservices failover?",
                "answer": (
                    "When our core payment API faced high latency during a Black Friday flash sale, "
                    "I had to prevent cascading timeouts across downstream services. I developed and deployed "
                    "a circuit breaker pattern with exponential backoff using Redis caching. As a result, we reduced "
                    "system latency by 45% and preserved 99.99% service availability with zero dropped transactions."
                ),
                "job_title": "Lead Backend Engineer"
            }
        )
        assert res.status_code == 200, f"Scoring failed at {endpoint}: {res.text}"
        data = res.json()
        assert 0.0 <= data["score"] <= 10.0
        assert len(data["feedback"]) > 10
        assert isinstance(data["strengths"], list)
        assert isinstance(data["improvements"], list)
    print("      PASSED: POST /interview/score returned score out of 10 and badge feedback!")

    # 8. Test Roadmap and checkbox persistence
    print("[5/6] Testing Roadmap 30/60/90 days and checkbox DB persistence...")
    res = requests.post(
        f"{BASE_URL}/api/generate/career-roadmap",
        headers=auth_headers,
        json={"analysis_id": analysis_id, "target_role": "Lead Backend Engineer"}
    )
    assert res.status_code == 201
    items = json.loads(res.json()["content"])
    assert len(items) >= 3

    for item in items:
        assert "skill" in item
        assert "why_it_matters" in item
        assert "mini_project" in item
        assert "phase" in item

    # Toggle checkbox
    toggle_res = requests.post(
        f"{BASE_URL}/api/analysis/{analysis_id}/roadmap-toggle",
        headers=auth_headers,
        json={"item_id": "item-1", "completed": True}
    )
    assert toggle_res.status_code == 200
    assert toggle_res.json()["roadmap_progress"]["item-1"] is True

    # Verify persistence by re-fetching analysis
    get_res = requests.get(f"{BASE_URL}/api/analysis/{analysis_id}", headers=auth_headers)
    assert get_res.status_code == 200
    saved_progress = get_res.json()["result_json"].get("roadmap_progress", {})
    assert saved_progress.get("item-1") is True
    print("      PASSED: Roadmap generated and checkbox state persisted in DB!")

    # 9. Test LinkedIn summary and History listing
    print("[6/6] Testing LinkedIn summary and History listing...")
    res = requests.post(
        f"{BASE_URL}/api/generate/linkedin-summary",
        headers=auth_headers,
        json={"analysis_id": analysis_id}
    )
    assert res.status_code == 201
    li_data = json.loads(res.json()["content"])
    assert len(li_data["headlines"]) == 3
    assert len(li_data["about"]) < 2000
    assert len(li_data["skills"]) == 15
    assert len(li_data["checklist"]) >= 5

    hist_res = requests.get(f"{BASE_URL}/api/analysis/?limit=50", headers=auth_headers)
    assert hist_res.status_code == 200
    analyses = hist_res.json()
    assert len(analyses) >= 1
    assert any(a["id"] == analysis_id for a in analyses)
    print("      PASSED: 3 headlines, <2000 char About, 15 skills, checklist & history listing!")

    print("\nALL PHASE 4 & PHASE 5 BACKEND TESTS PASSED SUCCESSFULLY! [6/6]")


if __name__ == "__main__":
    main()
